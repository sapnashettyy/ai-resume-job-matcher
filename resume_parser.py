
from pathlib import Path
import io

import pymupdf
from docx import Document


def extract_resume_text(file_name: str, file_bytes: bytes) -> str:
    """Extract readable text from an uploaded PDF or DOCX resume."""

    extension = Path(file_name).suffix.lower()

    if extension == ".pdf":
        with pymupdf.open(stream=file_bytes, filetype="pdf") as pdf:
            pages = [page.get_text() for page in pdf]
        text = "\n".join(pages)

    elif extension == ".docx":
        document = Document(io.BytesIO(file_bytes))

        paragraphs = [paragraph.text for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)

        text = "\n".join(paragraphs)

    else:
        raise ValueError("Please upload a PDF or DOCX file.")

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text found. The PDF may be scanned or image-based."
        )

    return text
